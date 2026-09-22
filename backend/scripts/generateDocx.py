import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, text_paragraphs, title="KEY TAKEAWAY", border_color="0F766E", bg_color="F0FDFA"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=150)
    
    # Left border only
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"★ {title}\n")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = RGBColor(15, 118, 110)
    
    for tp in text_paragraphs:
        p_body = cell.add_paragraph()
        p_body.paragraph_format.space_before = Pt(2)
        p_body.paragraph_format.space_after = Pt(3)
        run_body = p_body.add_run(tp)
        run_body.font.name = "Calibri"
        run_body.font.size = Pt(10)
        run_body.font.color.rgb = RGBColor(51, 65, 85)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def create_styled_document(output_path):
    doc = docx.Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles
    navy = RGBColor(27, 54, 93)     # #1B365D
    teal = RGBColor(15, 118, 110)   # #0F766E
    dark_gray = RGBColor(51, 65, 85)# #334155
    
    # Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = p_title.add_run("MadrasaHub School Onboarding &\nRole Customization Guide")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = navy
    
    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(2)
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("A Complete Operational Manual for Platform Admins and Self-Governing School Leaders")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = teal
    run_sub.font.italic = True
    
    # Horizontal Divider
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(14)
    p_div_run = p_div.add_run("―" * 48)
    p_div_run.font.color.rgb = RGBColor(203, 213, 225)
    
    # 1. Executive Summary
    h1 = doc.add_heading("1. Executive Overview & Design Philosophy", level=1)
    h1.style.font.name = "Arial"
    h1.style.font.color.rgb = navy
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    
    p = doc.add_paragraph(
        "MadrasaHub operates on a single core principle: The Platform Admin creates the school organization and hands over the primary Super Admin login. From that exact moment, the school is 100% self-governing and builds its own organizational empire."
    )
    p.style.font.name = "Calibri"
    p.style.font.size = Pt(11)
    p.style.font.color.rgb = dark_gray
    p.paragraph_format.space_after = Pt(8)
    
    p = doc.add_paragraph(
        "Schools come in all sizes. A large institution may employ dedicated staff for finance, human resources, admissions, and cashier booths. A growing community school may have a single Business Administrator handling billing, admissions, and fee collection at one desk, while the Principal directly runs payroll. MadrasaHub accommodates all of these structures natively without requiring custom code or developer intervention."
    )
    p.style.font.name = "Calibri"
    p.style.font.size = Pt(11)
    p.style.font.color.rgb = dark_gray
    p.paragraph_format.space_after = Pt(12)
    
    add_callout(
        doc,
        [
            "Zero code changes required: Schools rename roles, turn off unneeded roles, and reassign any of the 131 school-level features directly through the web UI.",
            "Dynamic menus: The sidebar adapts automatically so staff members see exactly the tools they have permission to use."
        ],
        title="CORE ARCHITECTURE PRINCIPLE"
    )

    # 2. Platform Admin's Role
    h2 = doc.add_heading("2. The Platform Admin's Role (Zero-Touch Setup)", level=1)
    h2.style.font.name = "Arial"
    h2.style.font.color.rgb = navy
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    
    p = doc.add_paragraph(
        "As the SaaS platform owner, your interaction with an individual school is clean, secure, and fast:"
    )
    p.style.font.name = "Calibri"
    p.style.font.size = Pt(11)
    p.style.font.color.rgb = dark_gray
    p.paragraph_format.space_after = Pt(6)
    
    steps = [
        ("Step 1: Create the School Tenant", "Open the Platform Admin portal, enter the school name, custom URL slug, contact details, and subscription plan tier."),
        ("Step 2: Generate the Super Admin Account", "Create the initial user with the role 'super_admin' and set a secure initial password."),
        ("Step 3: Hand Over the Credentials", "Deliver the login URL and credentials to the school principal or board. Your job is complete.")
    ]
    for step_title, step_desc in steps:
        p_step = doc.add_paragraph()
        p_step.paragraph_format.left_indent = Inches(0.2)
        p_step.paragraph_format.space_after = Pt(4)
        run_st = p_step.add_run(f"• {step_title}: ")
        run_st.bold = True
        run_st.font.name = "Calibri"
        run_st.font.size = Pt(10.5)
        run_st.font.color.rgb = navy
        run_sd = p_step.add_run(step_desc)
        run_sd.font.name = "Calibri"
        run_sd.font.size = Pt(10.5)
        run_sd.font.color.rgb = dark_gray

    # 3. Super Admin Capabilities
    h3 = doc.add_heading("3. What the School Super Admin Can Customize", level=1)
    h3.style.font.name = "Arial"
    h3.style.font.color.rgb = navy
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    
    features = [
        ("Rename Any Role", "Change standard keys to institutional titles (e.g., rename 'Finance Director' to 'Chief Financial Officer' or 'School Business Administrator')."),
        ("Turn Roles ON or OFF", "Toggle unused roles (such as Cashier or HR Manager) to Inactive. Inactive roles cannot be selected when creating new users."),
        ("Redistribute Any Feature", "Move any of the 131 school-level features to any active role with simple checkbox toggles in the Roles & Features matrix."),
        ("Dynamic Staff Menus", "Staff menus adapt in real time. If a user receives admissions permissions, the Admissions menu automatically appears in their navigation bar."),
        ("User Exceptions", "Grant special feature overrides or restrictions to an individual staff member without creating a whole new role.")
    ]
    for ft_title, ft_desc in features:
        p_ft = doc.add_paragraph()
        p_ft.paragraph_format.left_indent = Inches(0.2)
        p_ft.paragraph_format.space_after = Pt(4)
        run_ft = p_ft.add_run(f"✔ {ft_title}: ")
        run_ft.bold = True
        run_ft.font.name = "Calibri"
        run_ft.font.size = Pt(10.5)
        run_ft.font.color.rgb = teal
        run_fd = p_ft.add_run(ft_desc)
        run_fd.font.name = "Calibri"
        run_fd.font.size = Pt(10.5)
        run_fd.font.color.rgb = dark_gray

    # 4. Answers to Critical Operational Questions
    h4 = doc.add_heading("4. Common Operational Questions & Answers", level=1)
    h4.style.font.name = "Arial"
    h4.style.font.color.rgb = navy
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)

    qas = [
        ("Can a school give Registrar tasks to the Finance Director?", 
         "Yes. In Settings → Roles & Features, tick the Registrar features (e.g., Student Admissions, Student Directory, Class Enrollment) on the Finance role. The Finance Director's sidebar immediately shows both Finance and Registrar menus."),
        ("Can the school give HR and Payroll duties to the Super Admin?", 
         "Yes. The Super Admin simply ticks the HR and Payroll features on the Super Admin role. The HR Directory and Payroll Dashboard immediately appear in their management menu."),
        ("How does the system work without a dedicated Cashier?", 
         "The Super Admin turns the Cashier role OFF (inactive) and gives the Payment Collection feature to the Finance Director or front desk. Payments are collected directly at the Payment Desk without needing a dedicated cashier staff seat."),
        ("How does the Super Admin see which features are unassigned ('free')?", 
         "In the Roles & Features matrix (/tenant/roles), all 131 features are listed with checkboxes for each role. If a feature has zero checkboxes ticked across all active roles, it is unassigned and immediately visible.")
    ]
    for q, a in qas:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        r_q = p_q.add_run(f"Q: {q}")
        r_q.bold = True
        r_q.font.name = "Calibri"
        r_q.font.size = Pt(11)
        r_q.font.color.rgb = navy
        
        p_a = doc.add_paragraph()
        p_a.paragraph_format.left_indent = Inches(0.2)
        p_a.paragraph_format.space_after = Pt(6)
        r_a = p_a.add_run(f"A: {a}")
        r_a.font.name = "Calibri"
        r_a.font.size = Pt(10.5)
        r_a.font.color.rgb = dark_gray

    # 5. Disabling Staff & Handing Over Roles
    h5 = doc.add_heading("5. Disabling Staff & Safe Role Reassignment", level=1)
    h5.style.font.name = "Arial"
    h5.style.font.color.rgb = navy
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)
    
    p = doc.add_paragraph(
        "When an employee leaves the school or takes extended leave, their responsibilities can be transferred seamlessly:"
    )
    p.style.font.name = "Calibri"
    p.style.font.size = Pt(11)
    p.style.font.color.rgb = dark_gray
    p.paragraph_format.space_after = Pt(6)
    
    reassign_points = [
        ("Instant Account Deactivation", "In Staff → Users, setting an employee's status to Inactive immediately terminates their active sessions and blocks future logins."),
        ("Seamless Role Handover", "Because features belong to the Role rather than individual email addresses, assigning that role to a successor immediately equips them with all required menus and tools."),
        ("Accidental Lock-out Prevention", "MadrasaHub prevents removing role-management permissions from the last remaining active role that possesses it. School admins can never lock themselves out."),
        ("Active Role Protection", "A school cannot switch off a role while active users are currently assigned to it, ensuring no staff member is left in an undefined state.")
    ]
    for pt_t, pt_d in reassign_points:
        p_pt = doc.add_paragraph()
        p_pt.paragraph_format.left_indent = Inches(0.2)
        p_pt.paragraph_format.space_after = Pt(4)
        r_t = p_pt.add_run(f"• {pt_t}: ")
        r_t.bold = True
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(10.5)
        r_t.font.color.rgb = navy
        r_d = p_pt.add_run(pt_d)
        r_d.font.name = "Calibri"
        r_d.font.size = Pt(10.5)
        r_d.font.color.rgb = dark_gray

    # 6. Real-Life Operating Scenario
    h6 = doc.add_heading("6. Comprehensive Operational Scenario: 'Al-Hikmah Academy'", level=1)
    h6.style.font.name = "Arial"
    h6.style.font.color.rgb = navy
    h6.paragraph_format.space_before = Pt(14)
    h6.paragraph_format.space_after = Pt(6)
    
    p = doc.add_paragraph(
        "Al-Hikmah Academy is a growing school that operates with a highly efficient, compact administrative structure:"
    )
    p.style.font.name = "Calibri"
    p.style.font.size = Pt(11)
    p.style.font.color.rgb = dark_gray
    p.paragraph_format.space_after = Pt(6)

    # Table of Roles
    table = doc.add_table(rows=5, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    headers = ["Staff Role", "Assigned Responsibilities & Features", "Custom Sidebar Navigation"]
    col_widths = [Inches(1.8), Inches(2.6), Inches(2.1)]
    
    # Format Header Row
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "1B365D")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        p_hdr = hdr_cells[i].paragraphs[0]
        p_hdr.runs[0].font.bold = True
        p_hdr.runs[0].font.name = "Arial"
        p_hdr.runs[0].font.size = Pt(10)
        p_hdr.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        
    data = [
        ("Super Admin\n(Sheikh Omar)", 
         "• School Settings & Roles\n• Academic Years & Branches\n• HR Employee Directory\n• Payroll Review & Approval\n• Attendance Oversight",
         "• School Dashboard\n• School Settings\n• Staff & HR\n• Monthly Payroll\n• Academic Oversight"),
        ("School Business Admin\n(Sister Maryam)", 
         "• Fee Structures & Policy\n• Monthly Invoice Generation\n• Monthly Collection Sheet\n• Student Admissions & Directory\n• Class Enrollments & Transfers\n• Collect Payments (Oldest-First)\n• Receipts & Reversals",
         "• Admin Dashboard\n• Student Admissions\n• Student Directory\n• Invoices & Fees\n• Monthly Collection\n• Payments Desk"),
        ("Cashier", "Role Disabled (isActive = false)", "Hidden (No users assigned)"),
        ("HR Manager", "Role Disabled (isActive = false)", "Hidden (No users assigned)")
    ]
    
    for row_idx, row_data in enumerate(data, start=1):
        row_cells = table.rows[row_idx].cells
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_text in enumerate(row_data):
            row_cells[col_idx].text = cell_text
            row_cells[col_idx].width = col_widths[col_idx]
            set_cell_background(row_cells[col_idx], bg)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=100, right=100)
            p_cell = row_cells[col_idx].paragraphs[0]
            if len(p_cell.runs) > 0:
                p_cell.runs[0].font.name = "Calibri"
                p_cell.runs[0].font.size = Pt(9.5)
                p_cell.runs[0].font.color.rgb = dark_gray

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 7. Roadmap
    h7 = doc.add_heading("7. Step-by-Step Implementation Roadmap", level=1)
    h7.style.font.name = "Arial"
    h7.style.font.color.rgb = navy
    h7.paragraph_format.space_before = Pt(14)
    h7.paragraph_format.space_after = Pt(6)
    
    phases = [
        ("Phase 1: Platform Provisioning", [
            "Log into Platform Admin portal.",
            "Create School Tenant organization record.",
            "Generate primary Super Admin user credentials.",
            "Hand over access URL and temporary credentials to school owner."
        ]),
        ("Phase 2: Role & Feature Architecture", [
            "School Super Admin signs into tenant dashboard.",
            "Opens Settings → Roles & Features.",
            "Turns OFF unused roles (Cashier, HR Manager, etc.).",
            "Renames active roles to match institutional titles.",
            "Ticks/unticks features in the Matrix to distribute responsibilities."
        ]),
        ("Phase 3: Staff Account Creation", [
            "Opens Staff → Users.",
            "Creates staff accounts and selects their configured custom roles.",
            "Applies per-user permission adjustments if specific exceptions are needed."
        ]),
        ("Phase 4: Academic & Financial Launch", [
            "Set up Academic Year, Branches, and Classes.",
            "Define monthly fee structures per class or category.",
            "Admit and enroll active students.",
            "Execute monthly 1-click invoice generation.",
            "Collect student fees oldest-month first at the Payments Desk."
        ])
    ]
    
    for phase_title, phase_steps in phases:
        p_ph = doc.add_paragraph()
        p_ph.paragraph_format.space_before = Pt(6)
        p_ph.paragraph_format.space_after = Pt(2)
        r_ph = p_ph.add_run(f"■ {phase_title}")
        r_ph.bold = True
        r_ph.font.name = "Calibri"
        r_ph.font.size = Pt(11)
        r_ph.font.color.rgb = navy
        
        for s in phase_steps:
            p_s = doc.add_paragraph()
            p_s.paragraph_format.left_indent = Inches(0.25)
            p_s.paragraph_format.space_after = Pt(2)
            r_s = p_s.add_run(f"→ {s}")
            r_s.font.name = "Calibri"
            r_s.font.size = Pt(10)
            r_s.font.color.rgb = dark_gray
            
    doc.save(output_path)
    print(f"Successfully created: {output_path}")

if __name__ == "__main__":
    create_styled_document("c:\\Users\\ICT-LAB 3\\Documents\\school-management-system\\SCHOOL_ONBOARDING_AND_ROLE_CUSTOMIZATION.docx")
